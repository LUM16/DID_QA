"""Recommend DU teams for an uploaded delivery scope without training a model."""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from effort_prediction import (
    DEFAULT_SIMILARITY_CACHE_PATH,
    DEFAULT_SIMILARITY_CACHE_URL,
    TLF_SEMANTIC_MATCH_THRESHOLD,
    _cached_tlf_semantic_similarity,
    _jaccard,
    _load_similarity_cache,
    _normalized_name,
    _resolve_prediction_artifact,
    _save_similarity_cache,
    _similarity_candidates,
)
from neo4j_client import get_driver, load_env

TEAM_HISTORY_QUERY = """
MATCH (p:Person)-[:WORKS_ON]->(d:Delivery)
WHERE p.Team_Lead_Name IS NOT NULL
  AND trim(toString(p.Team_Lead_Name)) <> ''
  AND toLower(toString(d.DID_Status)) = 'completed'
  AND d.DID IS NOT NULL
  AND d.Actual_Delivery_Date IS NOT NULL
WITH DISTINCT p.Team_Lead_Name AS du_team, d
CALL (d) {
  OPTIONAL MATCH (d)-[t:HAS_TLF]->(tlf:TLF)
  RETURN collect(DISTINCT {
    name: tlf.Name, type: tlf.Type, source: tlf.Source,
    generation: t.Generation, qc: t.QC
  }) AS tlfs
}
CALL (d) {
  OPTIONAL MATCH (d)-[a:HAS_ADAM]->(adam:ADaM)
  RETURN collect(DISTINCT {
    name: adam.Name, generation: a.Generation, qc: a.QC
  }) AS adams
}
CALL (d) {
  OPTIONAL MATCH (d)-[s:HAS_SDTM]->(sdtm:SDTM)
  RETURN collect(DISTINCT {
    name: sdtm.Name, generation: s.Generation, qc: s.QC
  }) AS sdtms
}
RETURN du_team,
       toString(d.DID) AS did,
       substring(toString(d.Actual_Delivery_Date), 0, 10) AS completion_date,
       tlfs,
       adams,
       sdtms
"""

TEAM_WORKLOAD_QUERY = """
MATCH (p:Person)-[:WORKS_ON]->(d:Delivery)
WHERE p.Team_Lead_Name IS NOT NULL
  AND trim(toString(p.Team_Lead_Name)) <> ''
  AND toLower(toString(d.DID_Status)) IN ['ongoing', 'planned']
RETURN p.Team_Lead_Name AS du_team,
       count(DISTINCT d) AS active_did_count,
       count(DISTINCT p) AS active_people_count
"""

SEMANTIC_DID_CANDIDATE_LIMIT = 50


def _resolve_recommendation_cache_path(cache_path: Path) -> Path:
    """Resolve the default shared cache when Git-backed deployments receive an LFS pointer."""
    if cache_path.resolve() != DEFAULT_SIMILARITY_CACHE_PATH.resolve():
        return cache_path
    return _resolve_prediction_artifact(
        cache_path,
        DEFAULT_SIMILARITY_CACHE_PATH,
        "DID_EFFORT_SIMILARITY_CACHE_URL",
        DEFAULT_SIMILARITY_CACHE_URL,
    )


def _read_query(query: str) -> list[dict[str, Any]]:
    load_env()
    driver = get_driver()
    try:
        database = __import__("os").environ.get("NEO4J_DATABASE", "neo4j")
        with driver.session(database=database) as session:
            return session.execute_read(lambda tx: tx.run(query).data())
    finally:
        driver.close()


def _column(row: dict[str, Any], name: str) -> str:
    normalized = {
        "".join(character for character in key.lower() if character.isalnum()): value
        for key, value in row.items()
    }
    return str(normalized.get("".join(character for character in name.lower() if character.isalnum()), "") or "").strip()


def _split_sources(value: str) -> list[str]:
    return [item.strip() for item in value.replace(";", ",").split(",") if item.strip()]


def _scope_from_rows(
    tlf_rows: Iterable[dict[str, Any]], data_rows: Iterable[dict[str, Any]]
) -> dict[str, list[dict[str, str]]]:
    tlfs = []
    adams = []
    sdtms = []
    for row in tlf_rows:
        title = _column(row, "Title")
        if not title:
            continue
        source = _column(row, "Source Datasets")
        tlfs.append(
            {
                "name": title,
                "type": _column(row, "Type"),
                "source": source,
            }
        )
        for dataset in _split_sources(source):
            adams.append({"name": dataset})
    for row in data_rows:
        kind = _column(row, "SDTM/ADaM")
        name = _column(row, "Domain/Dataset Name")
        if not name:
            continue
        if _normalized_name(kind) == "ADAM":
            adams.append({"name": name})
        elif _normalized_name(kind) == "SDTM":
            sdtms.append({"name": name})
    return {
        "tlfs": _unique_items(tlfs),
        "adams": _unique_items(adams),
        "sdtms": _unique_items(sdtms),
    }


def _unique_items(items: Iterable[dict[str, Any]]) -> list[dict[str, str]]:
    unique: dict[tuple[str, str, str], dict[str, str]] = {}
    for item in items:
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        normalized = (
            _normalized_name(name),
            _normalized_name(item.get("type")),
            _normalized_name(item.get("source")),
        )
        unique[normalized] = {
            "name": name,
            "type": str(item.get("type") or "").strip(),
            "source": str(item.get("source") or "").strip(),
        }
    return list(unique.values())


def load_uploaded_scope(
    primary_file: Any, data_csv_file: Any | None = None
) -> dict[str, list[dict[str, str]]]:
    """Read an XLSX workbook with TLF/Data sheets or paired TLF/Data CSV files."""
    filename = str(getattr(primary_file, "name", "")).lower()
    content = primary_file.getvalue()
    if filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel upload support requires openpyxl.") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets = {sheet.title.strip().lower(): sheet for sheet in workbook.worksheets}
        if "tlf" not in sheets or "data" not in sheets:
            raise ValueError("The Excel workbook must contain sheets named TLF and Data.")

        def rows(sheet_name: str) -> list[dict[str, Any]]:
            values = sheets[sheet_name].iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values, ())]
            return [
                dict(zip(headers, values_row))
                for values_row in values
                if any(value is not None and str(value).strip() for value in values_row)
            ]

        return _scope_from_rows(rows("tlf"), rows("data"))
    if not filename.endswith(".csv") or data_csv_file is None:
        raise ValueError("Upload one .xlsx workbook, or both TLF CSV and Data CSV files.")
    tlf_rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    data_rows = list(
        csv.DictReader(io.StringIO(data_csv_file.getvalue().decode("utf-8-sig")))
    )
    return _scope_from_rows(tlf_rows, data_rows)


def _item_names(items: Iterable[dict[str, Any]]) -> set[str]:
    return {_normalized_name(item.get("name")) for item in items if item.get("name")}


def _date_or_minimum(value: Any) -> date:
    try:
        return datetime.fromisoformat(str(value)[:10]).date()
    except ValueError:
        return date.min


def _recency_score(completion_date: Any) -> float:
    days = (date.today() - _date_or_minimum(completion_date)).days
    if days <= 180:
        return 1.0
    if days <= 365:
        return 0.7
    if days <= 730:
        return 0.4
    return 0.15 if completion_date else 0.0


def recommend_teams(
    scope: dict[str, list[dict[str, str]]],
    history_rows: Iterable[dict[str, Any]],
    workload_rows: Iterable[dict[str, Any]],
    cache_path: Path = DEFAULT_SIMILARITY_CACHE_PATH,
    top_n: int = 3,
) -> dict[str, Any]:
    """Rank current DU teams from completed Delivery scope and active workload."""
    resolved_cache_path = _resolve_recommendation_cache_path(cache_path)
    target = {kind: _unique_items(scope.get(kind, [])) for kind in ("tlfs", "adams", "sdtms")}
    if not any(target.values()):
        raise ValueError("No TLF, ADaM, or SDTM scope was found in the uploaded file.")
    by_team: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in history_rows:
        team = str(row.get("du_team") or "").strip()
        did = str(row.get("did") or "").strip()
        if not team or not did:
            continue
        by_team[team][did] = {
            **row,
            "tlfs": _unique_items(row.get("tlfs") or []),
            "adams": _unique_items(row.get("adams") or []),
            "sdtms": _unique_items(row.get("sdtms") or []),
        }
    if not by_team:
        raise ValueError("Neo4j returned no completed DU team history.")

    workload = {
        str(row["du_team"]): int(row.get("active_did_count") or 0)
        for row in workload_rows
        if row.get("du_team")
    }
    maximum_workload = max(workload.values(), default=0)
    cache = _load_similarity_cache(resolved_cache_path)
    recommendations = []
    for team, histories in by_team.items():
        history = list(histories.values())
        team_items = {
            kind: _unique_items(
                item for record in history for item in record.get(kind, [])
            )
            for kind in ("tlfs", "adams", "sdtms")
        }
        coverage = {
            kind: (
                len(_item_names(target[kind]) & _item_names(team_items[kind]))
                / len(_item_names(target[kind]))
                if _item_names(target[kind])
                else 1.0
            )
            for kind in ("adams", "sdtms")
        }
        historical_similarity = []
        candidate_history = sorted(
            _similarity_candidates(target, history),
            key=lambda record: (
                _jaccard(
                    _item_names(target["adams"]), _item_names(record["adams"])
                )
                + _jaccard(
                    _item_names(target["sdtms"]), _item_names(record["sdtms"])
                ),
                _date_or_minimum(record.get("completion_date")),
            ),
            reverse=True,
        )[:SEMANTIC_DID_CANDIDATE_LIMIT]
        for record in candidate_history:
            tlf_similarity, tlf_coverage = _cached_tlf_semantic_similarity(
                target["tlfs"], record["tlfs"], cache
            )
            combined = (
                0.60 * tlf_coverage
                + 0.25 * _jaccard(_item_names(target["adams"]), _item_names(record["adams"]))
                + 0.15 * _jaccard(_item_names(target["sdtms"]), _item_names(record["sdtms"]))
            )
            historical_similarity.append(
                {
                    "did": record["did"],
                    "completion_date": record.get("completion_date"),
                    "combined_similarity": combined,
                    "tlf_semantic_coverage": tlf_coverage,
                    "tlf_semantic_similarity": tlf_similarity,
                }
            )
        historical_similarity.sort(
            key=lambda row: (row["combined_similarity"], _date_or_minimum(row["completion_date"])),
            reverse=True,
        )
        best = historical_similarity[0] if historical_similarity else None
        similar_70 = [
            row for row in historical_similarity
            if row["combined_similarity"] >= TLF_SEMANTIC_MATCH_THRESHOLD
        ]
        similar_85 = [
            row for row in historical_similarity if row["combined_similarity"] >= 0.85
        ]
        tlf_coverage = best["tlf_semantic_coverage"] if best else 0.0
        similar_score = (
            0.6 * (best["combined_similarity"] if best else 0.0)
            + 0.4 * min(1.0, len(similar_70) / 5)
        )
        recent_score = max(
            (_recency_score(row["completion_date"]) for row in similar_70),
            default=0.0,
        )
        active_dids = workload.get(team, 0)
        workload_score = (
            1.0 - active_dids / maximum_workload if maximum_workload else 1.0
        )
        score = (
            35 * tlf_coverage
            + 18 * coverage["adams"]
            + 12 * coverage["sdtms"]
            + 18 * similar_score
            + 10 * recent_score
            + 7 * workload_score
        )
        exact_tlf_names = _item_names(team_items["tlfs"])
        recommendations.append(
            {
                "du_team": team,
                "score": round(score, 1),
                "tlf_semantic_coverage": round(tlf_coverage, 3),
                "adam_coverage": round(coverage["adams"], 3),
                "sdtm_coverage": round(coverage["sdtms"], 3),
                "similar_did_count_ge_70": len(similar_70),
                "similar_did_count_ge_85": len(similar_85),
                "recent_experience_score": round(recent_score, 3),
                "active_did_count": active_dids,
                "completed_did_count": len(history),
                "semantic_candidate_did_count": len(candidate_history),
                "missing_adams": sorted(
                    item["name"] for item in target["adams"]
                    if _normalized_name(item["name"]) not in _item_names(team_items["adams"])
                ),
                "missing_sdtms": sorted(
                    item["name"] for item in target["sdtms"]
                    if _normalized_name(item["name"]) not in _item_names(team_items["sdtms"])
                ),
                "missing_exact_tlfs": sorted(
                    item["name"] for item in target["tlfs"]
                    if _normalized_name(item["name"]) not in exact_tlf_names
                ),
                "similar_dids": historical_similarity[:5],
            }
        )
    recommendations.sort(key=lambda row: (-row["score"], row["du_team"]))
    for rank, row in enumerate(recommendations, start=1):
        row["rank"] = rank
        row["recommendation_level"] = (
            "Recommended" if row["score"] >= 80
            and not row["missing_adams"] and not row["missing_sdtms"]
            else "Suitable with review" if row["score"] >= 65
            else "Backup option" if row["score"] >= 45
            else "Insufficient evidence"
        )
    _save_similarity_cache(cache, resolved_cache_path)
    return {
        "target_counts": {kind: len(target[kind]) for kind in target},
        "recommendations": recommendations[:top_n],
        "cache": {
            "hits": cache["hits"],
            "misses": cache["misses"],
            "candidate_pairs": cache["hits"] + cache["misses"],
        },
    }


def recommend_uploaded_scope(primary_file: Any, data_csv_file: Any | None = None) -> dict[str, Any]:
    """Load an uploaded scope and retrieve current Neo4j evidence for ranking."""
    return recommend_teams(
        load_uploaded_scope(primary_file, data_csv_file),
        _read_query(TEAM_HISTORY_QUERY),
        _read_query(TEAM_WORKLOAD_QUERY),
    )
